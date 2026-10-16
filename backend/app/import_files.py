"""Deterministic file inspection. Cell contents never become instructions or code."""
import base64
import binascii
import csv
import hashlib
import io
import re
import time
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Literal

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field

from app.sql_executor import QueryError

MAX_FILE = 5 * 1024 * 1024
MAX_ROWS, MAX_COLUMNS = 10000, 50
FIELDS = ('sale_date', 'product_name', 'quantity', 'unit', 'unit_price', 'amount', 'currency')
REQUIRED = ('sale_date', 'product_name', 'quantity', 'amount')
ALIASES = {
    'sale_date': ('date', 'sale date', 'sale_date', 'tarikh'),
    'product_name': ('product', 'product name', 'product_name', 'item'),
    'quantity': ('quantity', 'qty'), 'unit': ('unit', 'units'),
    'unit_price': ('unit price', 'unit_price', 'rate'),
    'amount': ('amount', 'total', 'line total', 'sales'), 'currency': ('currency',),
}
UNITS = {'litre': 'litre', 'liter': 'litre', 'litres': 'litre', 'liters': 'litre',
         'l': 'litre', 'kg': 'kg', 'kilogram': 'kg', 'piece': 'piece', 'pieces': 'piece', 'pcs': 'piece'}


class FileRequest(BaseModel):
    model_config = ConfigDict(extra='forbid')
    filename: str = Field(min_length=1, max_length=200)
    content: str = Field(min_length=1, max_length=6990508)
    sheet: str | None = Field(default=None, max_length=100)


class ValidationRequest(FileRequest):
    mapping: dict[str, str] = Field(max_length=7)
    date_format: Literal['ISO', 'DMY', 'MDY'] = 'ISO'
    default_unit: Literal['litre', 'kg', 'piece'] | None = None
    currency: Literal['INR'] = 'INR'
    amount_confirmed: bool = False
    adjustment_note: str = Field(default='', max_length=300)


def decode_file(body: FileRequest) -> bytes:
    try:
        raw = base64.b64decode(body.content, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise QueryError(422, 'File encoding is invalid. Choose the file again.') from exc
    if not raw or len(raw) > MAX_FILE:
        raise QueryError(413, 'Choose a nonempty file up to 5 MiB.')
    return raw


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    result = str(value).strip()
    if len(result) > 1000 or '\x00' in result:
        raise QueryError(422, 'Cells must contain at most 1000 characters and no NUL bytes.')
    return result


def read_file(body: FileRequest):
    raw = decode_file(body)
    started = time.monotonic()
    workbook = None
    sheets = ['CSV']
    try:
        if body.filename.lower().endswith('.csv'):
            source = raw.decode('utf-8-sig')
            if '\x00' in source:
                raise ValueError('binary CSV')
            rows = csv.reader(io.StringIO(source, newline=''), strict=True)
            if body.sheet not in (None, 'CSV'):
                raise QueryError(422, 'Choose CSV as the sheet.')
        elif body.filename.lower().endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries = archive.infolist()
                if (len(entries) > 200 or sum(entry.file_size for entry in entries) > 20 * 1024 * 1024
                        or any(entry.flag_bits & 1 for entry in entries)):
                    raise QueryError(422, 'Workbook exceeds expanded-size limits or is encrypted.')
                if any('vbaProject' in entry.filename for entry in entries):
                    raise QueryError(422, 'Macro-enabled workbooks are unsupported.')
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            sheets = workbook.sheetnames
            if not sheets or len(sheets) > 10:
                raise QueryError(422, 'Workbook must have 1–10 sheets.')
            if body.sheet is None and len(sheets) > 1:
                return raw, sheets, None, [], []
            selected = body.sheet or sheets[0]
            if selected not in sheets:
                raise QueryError(422, 'Choose an available sheet.')
            sheet = workbook[selected]
            # Do not trust dimensions supplied by an uploaded workbook.
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
        else:
            raise QueryError(422, 'Only UTF-8 comma-separated .csv and .xlsx files are supported.')
        headers = None
        data = []
        for source_row, row in enumerate(rows, 1):
            if time.monotonic() - started > 5:
                raise QueryError(422, 'File parsing budget exceeded. Split the file.')
            if source_row > MAX_ROWS + 1 or len(row) > MAX_COLUMNS:
                raise QueryError(422, 'Limit: 10000 data rows and 50 columns per sheet.')
            values = [cell_text(value) for value in row]
            if headers is None:
                headers = values
                if not headers or any(not value for value in headers) or len(set(headers)) != len(headers):
                    raise QueryError(422, 'First row must contain unique, nonempty column names.')
                continue
            if len(values) > len(headers):
                raise QueryError(422, f'Row {source_row} has more cells than the header.')
            values += [''] * (len(headers) - len(values))
            # Blank records remain visible and fail validation rather than disappearing.
            data.append((source_row, values))
        if not headers or not data:
            raise QueryError(422, 'File needs a header and at least one data row.')
        return raw, sheets, body.sheet or sheets[0], headers, data
    except QueryError:
        raise
    except Exception as exc:
        raise QueryError(422, 'Cannot read file. Use a valid UTF-8 CSV or unencrypted XLSX workbook.') from exc
    finally:
        if workbook:
            workbook.close()


def preview_file(body: FileRequest):
    raw, sheets, sheet, headers, data = read_file(body)
    proposed = {}
    for field, aliases in ALIASES.items():
        matches = [header for header in headers if header.lower() in aliases]
        if len(matches) == 1:
            proposed[field] = matches[0]
    types = {}
    for index, header in enumerate(headers):
        values = [row[index] for _, row in data[:20] if row[index]]
        types[header] = ('number' if values and all(re.fullmatch(r'-?\d+(\.\d+)?', v) for v in values)
                         else 'date' if values and all(re.fullmatch(r'\d{4}-\d{2}-\d{2}.*', v) for v in values)
                         else 'text / uncertain')
    return {'sha256': hashlib.sha256(raw).hexdigest(), 'sheets': sheets, 'sheet': sheet,
            'requires_sheet': sheet is None, 'headers': headers, 'preview': [row for _, row in data[:10]],
            'row_count': len(data), 'proposed_mapping': proposed, 'types': types,
            'required_fields': list(REQUIRED) + ['unit or default_unit'],
            'notice': 'Preview only. Confirm mappings, date order, unit, and amount meaning before import.'}


def number(value, scale, positive=False):
    if not re.fullmatch(r'\d+(\.\d+)?', value):
        raise ValueError('Use a nonnegative decimal with dot separator; no commas, symbols or formulas.')
    try:
        result = Decimal(value)
        if result >= Decimal(10) ** (18 - scale) or result.as_tuple().exponent < -scale:
            raise ValueError(f'Number exceeds supported precision ({scale} decimal places).')
        if positive and result <= 0:
            raise ValueError('Quantity must be greater than zero.')
        return result
    except InvalidOperation as exc:
        raise ValueError('Invalid decimal number.') from exc


def validate_file(body: ValidationRequest):
    raw, _, sheet, headers, data = read_file(body)
    if sheet is None:
        raise QueryError(422, 'Choose a sheet before validating.')
    mapping = body.mapping
    if (set(mapping) - set(FIELDS) or any(field not in mapping for field in REQUIRED)
            or ('unit' not in mapping and not body.default_unit)
            or any(value not in headers for value in mapping.values())
            or len(set(mapping.values())) != len(mapping)):
        raise QueryError(422, 'Map each required field to a distinct column; choose a unit column or default unit.')
    errors, normalized, diagnostics = [], [], []
    invalid_count = 0
    for row_num, row in data:
        values = {field: row[headers.index(header)] for field, header in mapping.items()}
        try:
            raw_date = values['sale_date']
            if re.fullmatch(r'\d{4}-\d{2}-\d{2}(T00:00:00)?', raw_date):
                sale_date = date.fromisoformat(raw_date[:10])
            else:
                if body.date_format == 'ISO':
                    raise ValueError('sale_date: use YYYY-MM-DD, or explicitly select DMY / MDY.')
                pattern = '%d/%m/%Y' if body.date_format == 'DMY' else '%m/%d/%Y'
                sale_date = datetime.strptime(raw_date, pattern).date()
            product = values['product_name']
            if not product or len(product) > 300:
                raise ValueError('product_name: enter 1–300 characters.')
            if product.startswith(('=', '+', '-', '@')):
                raise ValueError('product_name: formula-like text is unsupported; supply a plain product name.')
            quantity = number(values['quantity'], 3, positive=True)
            amount = number(values['amount'], 2)
            price = number(values['unit_price'], 2) if values.get('unit_price') else None
            unit = UNITS.get(values.get('unit', body.default_unit or '').lower())
            if unit is None:
                raise ValueError('unit: use litre, kg or piece; unknown units need correction.')
            if 'currency' in values and values['currency'] != body.currency:
                raise ValueError('currency: only INR is supported; mixed currencies need separate conversion outside this app.')
            if price is not None and abs(quantity * price - amount) > Decimal('0.01'):
                diagnostics.append({'row': row_num, 'message': 'quantity × unit_price differs from amount. Confirm tax/discount/rounding meaning.'})
            normalized.append({'source_row': row_num, 'sale_date': sale_date, 'product_name': product,
                               'quantity': quantity, 'unit': unit, 'unit_price': price, 'amount': amount})
        except ValueError as exc:
            invalid_count += 1
            if len(errors) < 100:
                errors.append({'row': row_num, 'message': str(exc)})
    confirmations = []
    if not body.amount_confirmed:
        confirmations.append('Confirm amount is final line revenue, including documented tax/discounts; revenue is not profit.')
    if diagnostics and not body.adjustment_note.strip():
        confirmations.append('Explain amount differences in the adjustment note before importing.')
    report = {'valid': not invalid_count and not confirmations, 'accepted_rows': len(normalized),
              'rejected_rows': invalid_count, 'errors': errors, 'errors_truncated': invalid_count > len(errors),
              'diagnostics': diagnostics[:100], 'diagnostic_count': len(diagnostics),
              'confirmations': confirmations, 'sheet': sheet, 'currency': body.currency,
              'sha256': hashlib.sha256(raw).hexdigest(), 'committed': False}
    return report, normalized
